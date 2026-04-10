"""FastAPI backend for Easy AI — fully integrated multi-module AI workspace."""

from __future__ import annotations

import io
import json
import os
import secrets
from datetime import datetime, timezone
from pathlib import Path

from fastapi import Depends, FastAPI, File, Form, HTTPException, UploadFile, status
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, HTMLResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, EmailStr
from sqlalchemy.orm import Session

from main_ai    import MainAI, detect_module, _MODULE_NAMES
from education_orchestrator import EducationOrchestrator, MODULE_INFO
from export_utils       import export_pdf, export_docx
from advanced_export    import export_html, export_markdown, export_txt, export_json, export_zip
from workspace_store    import save_workspace_item, list_workspace_items
from education_memory   import add_chat_message, get_chat_history
from uploaded_material_store import save_uploaded_material, list_uploaded_materials
from upload_ai          import UploadAI
from citation_ai        import CitationAI
from search_ai          import get_search_engine
from workflow_ai        import run_workflow, WORKFLOWS
from database           import (
    get_db, create_user, get_user_by_email, save_document, list_documents,
    search_documents, create_api_key, get_api_key_owner, check_rate_limit,
    increment_gen_count, save_user_template, list_user_templates,
    get_document_by_share_token, User, Document, ApiKey,
)
from auth import (
    hash_password, verify_password, create_access_token,
    get_current_user, get_current_user_optional,
)

app = FastAPI(
    title="Easy AI",
    description="Full multi-module AI workspace — 17 AI agents, real file generation, global search",
    version="2.0.0",
)

_STATIC_DIR = Path(__file__).parent / "static"
_STATIC_DIR.mkdir(exist_ok=True)
_DIST_DIR   = Path(__file__).parent.parent / "frontend" / "dist"

app.mount("/static", StaticFiles(directory=str(_STATIC_DIR)), name="static")
if (_DIST_DIR / "assets").exists():
    app.mount("/assets", StaticFiles(directory=str(_DIST_DIR / "assets")), name="assets")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)

_brain    = MainAI()
_edu      = EducationOrchestrator()
_upload   = UploadAI()
_citation = CitationAI()
_search   = get_search_engine()

_MIME = {
    ".pdf":  "application/pdf",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".html": "text/html",
    ".md":   "text/markdown",
    ".txt":  "text/plain",
    ".json": "application/json",
    ".zip":  "application/zip",
}

# Pre-built templates
TEMPLATES: list[dict] = [
    {"id":"doc-proposal",  "module":"document",   "label":"Business Proposal",          "prompt":"Write a professional business proposal for an AI-powered education platform"},
    {"id":"doc-report",    "module":"document",   "label":"Executive Report",            "prompt":"Create an executive report on the current state of remote work in tech companies"},
    {"id":"doc-contract",  "module":"document",   "label":"Service Contract",            "prompt":"Draft a service agreement contract for a software development project"},
    {"id":"doc-resume",    "module":"document",   "label":"Professional Resume",         "prompt":"Create a professional resume for a senior software engineer with 10 years experience"},
    {"id":"biz-plan",      "module":"business",   "label":"Full Business Plan",          "prompt":"Create a complete business plan for a SaaS startup in education technology"},
    {"id":"biz-pitch",     "module":"business",   "label":"Investor Pitch Deck",         "prompt":"Write an investor pitch deck for a Series A fundraise for an AI platform startup"},
    {"id":"biz-swot",      "module":"business",   "label":"SWOT Analysis",               "prompt":"Conduct a SWOT analysis for a mid-sized consulting firm entering the AI market"},
    {"id":"biz-marketing", "module":"business",   "label":"Marketing Strategy",          "prompt":"Create a digital marketing strategy for a B2B SaaS product targeting HR professionals"},
    {"id":"res-litreview", "module":"research",   "label":"Literature Review",           "prompt":"Write a literature review on the impact of artificial intelligence on higher education"},
    {"id":"res-paper",     "module":"research",   "label":"Research Paper",              "prompt":"Write a full research paper on the effectiveness of online vs classroom learning"},
    {"id":"ana-kpi",       "module":"analytics",  "label":"KPI Dashboard Report",        "prompt":"Create a KPI dashboard report for a customer success team — retention and NPS"},
    {"id":"con-blog",      "module":"content",    "label":"SEO Blog Post",               "prompt":"Write an SEO-optimized blog post about the top 10 benefits of AI for small businesses"},
    {"id":"con-social",    "module":"content",    "label":"Social Media Pack",           "prompt":"Create a social media content pack for LinkedIn, Twitter, and Instagram on productivity"},
    {"id":"edu-lesson",    "module":"teacher",    "label":"Lesson Plan",                 "prompt":"Create a detailed lesson plan on photosynthesis for 8th grade biology students"},
    {"id":"edu-syllabus",  "module":"professor",  "label":"University Syllabus",         "prompt":"Design a complete university syllabus for an undergraduate course in Organisational Behaviour"},
    {"id":"edu-quiz",      "module":"exam",       "label":"Quiz with Answer Key",        "prompt":"Create a 10-question quiz with answer key on the French Revolution for high school students"},
    {"id":"edu-study",     "module":"student",    "label":"Study Guide",                 "prompt":"Create a comprehensive study guide with flashcards for Newton's Laws of Motion"},
    {"id":"edu-course",    "module":"course",     "label":"Course Curriculum",           "prompt":"Build a 10-week Data Science curriculum for beginners with no programming experience"},
    {"id":"leg-case",      "module":"court",      "label":"Arbitration Case Package",    "prompt":"Prepare a full legal case package for an employment discrimination arbitration case"},
    {"id":"leg-contract",  "module":"court",      "label":"Commercial Contract",         "prompt":"Draft a comprehensive commercial contract between a digital agency and its client"},
    {"id":"sim-interview", "module":"simulation", "label":"Job Interview Prep",          "prompt":"Create a mock job interview simulation for a Senior Product Manager role"},
    {"id":"xls-budget",    "module":"excel",      "label":"Annual Budget",               "prompt":"Create an annual budget spreadsheet for a tech startup with 20 employees"},
    {"id":"ppt-business",  "module":"slides",     "label":"Business Presentation",       "prompt":"Create an 8-slide business presentation on AI in healthcare"},
    {"id":"cit-apa",       "module":"citation",   "label":"APA Citation Guide",          "prompt":"Generate APA citation examples for a research paper on climate change"},
    {"id":"cit-oscola",    "module":"citation",   "label":"OSCOLA Legal Citations",      "prompt":"Generate OSCOLA citation examples for an employment law case"},
]

_MODULE_ICON = {
    "document":"📄","slides":"📽️","excel":"📊","professor":"🎓","teacher":"📚",
    "exam":"📝","simulation":"💼","court":"⚖️","student":"🙋","admin":"🗂️",
    "multilingual":"🌐","integrity":"🛡️","business":"🏢","research":"🔬",
    "analytics":"📈","content":"✍️","course":"🏫","upload":"��","citation":"📖",
}


# ── Helpers ────────────────────────────────────────────────────────────────

def _url(path: str | None) -> str | None:
    return f"/download/{Path(path).name}" if path else None


def _save_to_db_and_search(
    db, user_id, prompt, module, title, body, sections, raw
) -> None:
    try:
        doc = save_document(
            db, user_id=user_id, title=title, module=module,
            prompt=prompt, body=body,
            sections_json=json.dumps(sections),
            pdf_url=_url(raw.get("pdf_path")),
            docx_url=_url(raw.get("docx_path")),
            pptx_url=_url(raw.get("pptx_path")),
            xlsx_url=_url(raw.get("xlsx_url")),
        )
        # Also index in in-memory search
        _search.add_document({
            "id":       doc.id,
            "user_id":  user_id,
            "title":    title,
            "module":   module,
            "prompt":   prompt,
            "body":     body,
            "pdf_url":  _url(raw.get("pdf_path")),
            "docx_url": _url(raw.get("docx_path")),
            "pptx_url": _url(raw.get("pptx_path")),
            "xlsx_url": _url(raw.get("xlsx_url")),
            "created_at": datetime.now(timezone.utc).isoformat(),
        })
        # Legacy workspace file store
        save_workspace_item({
            "prompt": prompt, "mode": module, "module": module,
            "title": title, "preview": body[:300],
            "pdf_url":  _url(raw.get("pdf_path")),
            "docx_url": _url(raw.get("docx_path")),
            "pptx_url": _url(raw.get("pptx_path")),
            "xlsx_url": _url(raw.get("xlsx_url")),
            "timestamp": datetime.now(timezone.utc).isoformat(),
        })
    except Exception:
        pass  # Never crash the main request


def _extract_text(raw_bytes: bytes, filename: str) -> str:
    fname = (filename or "").lower()
    if any(fname.endswith(e) for e in (".txt", ".md", ".csv", ".json", ".xml", ".html")):
        return raw_bytes.decode("utf-8", errors="replace")
    if fname.endswith(".pdf"):
        try:
            import PyPDF2
            reader = PyPDF2.PdfReader(io.BytesIO(raw_bytes))
            return "\n".join((p.extract_text() or "") for p in reader.pages)
        except Exception:
            return raw_bytes.decode("utf-8", errors="replace")
    if fname.endswith(".docx"):
        try:
            from docx import Document as DocxDoc
            doc = DocxDoc(io.BytesIO(raw_bytes))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        except Exception:
            return raw_bytes.decode("utf-8", errors="replace")
    if fname.endswith((".xlsx", ".xls")):
        try:
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(raw_bytes), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    line = " | ".join(str(c) for c in row if c is not None)
                    if line.strip():
                        lines.append(line)
            return "\n".join(lines)
        except Exception:
            return raw_bytes.decode("utf-8", errors="replace")
    return raw_bytes.decode("utf-8", errors="replace")


# ── SPA ────────────────────────────────────────────────────────────────────

@app.get("/", response_class=HTMLResponse, response_model=None, include_in_schema=False)
def home():
    index = _DIST_DIR / "index.html"
    if index.is_file():
        return FileResponse(str(index))
    return HTMLResponse("<h1>Easy AI</h1><p>Frontend not built. Run: <code>cd frontend && npm run build</code></p>", status_code=503)

@app.get("/health")
def health() -> dict:
    return {"status": "ok", "version": "2.0.0", "modules": len(_MODULE_NAMES), "search_index": _search.size()}


# ── AUTH ───────────────────────────────────────────────────────────────────

class SignUpRequest(BaseModel):
    email: str
    password: str
    name: str
    role: str = "student"

class LoginRequest(BaseModel):
    email: str
    password: str

class ProfileUpdate(BaseModel):
    name: str | None = None
    role: str | None = None
    avatar_url: str | None = None


@app.post("/auth/signup")
def signup(req: SignUpRequest, db: Session = Depends(get_db)) -> dict:
    if get_user_by_email(db, req.email):
        raise HTTPException(status_code=400, detail="Email already registered.")
    if len(req.password) < 8:
        raise HTTPException(status_code=400, detail="Password must be at least 8 characters.")
    user = create_user(db, email=req.email, name=req.name,
                       hashed_pw=hash_password(req.password), role=req.role)
    token = create_access_token(user.id, user.email, user.role, user.plan)
    return {"token": token, "user": _user_dict(user)}


@app.post("/auth/login")
def login(req: LoginRequest, db: Session = Depends(get_db)) -> dict:
    user = get_user_by_email(db, req.email)
    if not user or not verify_password(req.password, user.hashed_pw):
        raise HTTPException(status_code=401, detail="Invalid email or password.")
    token = create_access_token(user.id, user.email, user.role, user.plan)
    return {"token": token, "user": _user_dict(user)}


@app.get("/auth/me")
def me(user: User = Depends(get_current_user)) -> dict:
    return _user_dict(user)


@app.patch("/auth/profile")
def update_profile(req: ProfileUpdate,
                   user: User = Depends(get_current_user),
                   db: Session = Depends(get_db)) -> dict:
    if req.name:       user.name = req.name
    if req.role:       user.role = req.role
    if req.avatar_url: user.avatar_url = req.avatar_url
    db.commit()
    db.refresh(user)
    return _user_dict(user)


def _user_dict(u: User) -> dict:
    return {
        "id": u.id, "email": u.email, "name": u.name, "role": u.role,
        "plan": u.plan, "avatar_url": u.avatar_url,
        "gen_count": u.gen_count, "gen_date": u.gen_date,
    }


# ── CENTRAL BRAIN ─────────────────────────────────────────────────────────

class BuildRequest(BaseModel):
    prompt: str
    module: str = "auto"

class SimpleGenerateRequest(BaseModel):
    prompt: str


@app.post("/api/build")
def api_build(
    req: BuildRequest,
    db: Session = Depends(get_db),
    user: User | None = Depends(get_current_user_optional),
) -> dict:
    if not req.prompt.strip():
        raise HTTPException(status_code=422, detail="Prompt must not be empty.")

    # Check rate limit for authenticated users
    if user:
        allowed, remaining = check_rate_limit(db, user)
        if not allowed:
            raise HTTPException(
                status_code=429,
                detail=f"Daily generation limit reached for your {user.plan} plan. Upgrade to Pro for unlimited access.",
            )

    try:
        raw = _brain.run(req.prompt.strip(), module_hint=req.module)
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Generation failed: {e}")

    module   = raw.get("module", "document")
    title    = raw["title"]
    body     = raw["body"]
    sections = raw["sections"]
    user_id  = user.id if user else None

    if user:
        increment_gen_count(db, user)

    _save_to_db_and_search(db, user_id, req.prompt, module, title, body, sections, raw)

    return {
        "success":    True,
        "module":     module,
        "module_name":raw.get("module_name", _MODULE_NAMES.get(module, "Easy AI")),
        "title":      title,
        "body":       body,
        "sections":   sections,
        "pdf_url":    _url(raw.get("pdf_path")),
        "docx_url":   _url(raw.get("docx_path")),
        "pptx_url":   _url(raw.get("pptx_path")),
        "xlsx_url":   _url(raw.get("xlsx_path")),
        "meta": {
            "intent":    module,
            "timestamp": datetime.now(timezone.utc).isoformat(),
            "remaining_today": (check_rate_limit(db, user)[1] if user else None),
        },
    }


@app.post("/api/detect-module")
def api_detect_module(req: SimpleGenerateRequest) -> dict:
    module = detect_module(req.prompt)
    return {"module": module, "module_name": _MODULE_NAMES.get(module, "Easy AI")}


# ── LEGACY /generate ──────────────────────────────────────────────────────

@app.post("/generate")
async def generate(req: SimpleGenerateRequest, db: Session = Depends(get_db)) -> dict:
    try:
        result = _brain.run(req.prompt, doc_type="auto")
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Generation failed: {exc}")
    return {
        "content": result["body"],
        "pdf":  _url(result.get("pdf_path")),
        "docx": _url(result.get("docx_path")),
    }


# ── DOWNLOAD ──────────────────────────────────────────────────────────────

@app.get("/download/{filename}")
def download_file(filename: str) -> FileResponse:
    import tempfile
    safe_name  = os.path.basename(filename)
    export_dir = os.path.realpath(os.path.join(tempfile.gettempdir(), "docgen_exports"))
    resolved   = os.path.realpath(os.path.join(export_dir, safe_name))
    if not resolved.startswith(export_dir + os.sep):
        raise HTTPException(status_code=403, detail="Access denied.")
    if not os.path.isfile(resolved):
        raise HTTPException(status_code=404, detail="File not found.")
    ext = Path(resolved).suffix.lower()
    return FileResponse(resolved, media_type=_MIME.get(ext, "application/octet-stream"), filename=safe_name)


# ── ADVANCED EXPORTS ──────────────────────────────────────────────────────

class ExportRequest(BaseModel):
    title: str
    sections: list[dict]
    format: str  # html | md | txt | json

@app.post("/api/export")
def api_export(req: ExportRequest) -> dict:
    """Export existing sections to a new format."""
    fmt = req.format.lower()
    if fmt == "html":
        path = export_html(req.title, req.sections)
    elif fmt in ("md", "markdown"):
        path = export_markdown(req.title, req.sections)
    elif fmt == "txt":
        path = export_txt(req.title, req.sections)
    elif fmt == "json":
        path = export_json(req.title, req.sections)
    else:
        raise HTTPException(status_code=400, detail=f"Unsupported format: {fmt}. Use html, md, txt, json.")
    return {"url": _url(path), "format": fmt}


class BatchExportRequest(BaseModel):
    document_ids: list[str]

@app.post("/api/export/batch")
def api_batch_export(
    req: BatchExportRequest,
    db: Session = Depends(get_db),
    user: User | None = Depends(get_current_user_optional),
) -> dict:
    """Batch export multiple documents as a ZIP archive."""
    docs_data = []
    for doc_id in req.document_ids[:20]:  # cap at 20
        doc = db.query(Document).filter(Document.id == doc_id, Document.is_deleted == False).first()
        if doc:
            secs = json.loads(doc.sections_json) if doc.sections_json else []
            docs_data.append({"title": doc.title, "sections": secs})
    if not docs_data:
        raise HTTPException(status_code=404, detail="No documents found.")
    zip_path = export_zip(docs_data)
    return {"url": _url(zip_path), "count": len(docs_data)}


# ── CITATION ──────────────────────────────────────────────────────────────

class CitationRequest(BaseModel):
    prompt: str

@app.post("/api/citation")
def api_citation(req: CitationRequest, db: Session = Depends(get_db),
                 user: User | None = Depends(get_current_user_optional)) -> dict:
    if not req.prompt.strip():
        raise HTTPException(status_code=422, detail="Prompt must not be empty.")
    result = _citation.generate(req.prompt.strip())
    pdf_path  = export_pdf(result["title"],  result["sections"])
    docx_path = export_docx(result["title"], result["sections"])
    _save_to_db_and_search(db, user.id if user else None, req.prompt, "citation",
                           result["title"], result["body"], result["sections"],
                           {"pdf_path": pdf_path, "docx_path": docx_path})
    return {
        "success":  True,
        "module":   "citation",
        "title":    result["title"],
        "body":     result["body"],
        "sections": result["sections"],
        "pdf_url":  _url(pdf_path),
        "docx_url": _url(docx_path),
    }


# ── SEARCH ────────────────────────────────────────────────────────────────

class SearchRequest(BaseModel):
    query: str
    module: str | None = None
    limit: int = 20

@app.post("/api/search")
def api_search(
    req: SearchRequest,
    db: Session = Depends(get_db),
    user: User | None = Depends(get_current_user_optional),
) -> dict:
    if not req.query.strip():
        return {"results": [], "count": 0}
    user_id = user.id if user else None
    # Try DB search first (persisted)
    db_results = search_documents(db, req.query, user_id=user_id, module=req.module, limit=req.limit)
    # Supplement with in-memory search engine (catches current session docs)
    mem_results = _search.search(req.query, user_id=user_id, module=req.module, limit=req.limit)
    # Merge, dedup by id
    seen: set[str] = {r["id"] for r in db_results}
    merged = db_results[:]
    for r in mem_results:
        if r["id"] not in seen:
            merged.append(r)
            seen.add(r["id"])
    merged.sort(key=lambda x: x.get("score", 0), reverse=True)
    return {"results": merged[:req.limit], "count": len(merged)}


@app.get("/api/search/suggestions")
def api_search_suggestions(q: str = "", limit: int = 8) -> list[str]:
    if not q.strip():
        return []
    return _search.suggestions(q.strip(), limit=limit)


# ── WORKFLOW ──────────────────────────────────────────────────────────────

class WorkflowRequest(BaseModel):
    action: str
    source_content: str
    source_title: str
    params: dict = {}

@app.get("/api/workflows")
def list_workflows() -> list:
    return [{"id": k, **v} for k, v in WORKFLOWS.items()]

@app.post("/api/workflow/{action}")
def run_workflow_endpoint(
    action: str,
    req: WorkflowRequest,
    db: Session = Depends(get_db),
    user: User | None = Depends(get_current_user_optional),
) -> dict:
    try:
        raw = run_workflow(action, req.source_content, req.source_title, req.params)
    except ValueError as e:
        raise HTTPException(status_code=400, detail=str(e))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Workflow failed: {e}")
    module   = raw.get("module", "document")
    user_id  = user.id if user else None
    _save_to_db_and_search(db, user_id, f"[Workflow:{action}] {req.source_title}",
                           module, raw["title"], raw["body"], raw["sections"], raw)
    return {
        "success":  True,
        "action":   action,
        "module":   module,
        "title":    raw["title"],
        "body":     raw["body"],
        "sections": raw["sections"],
        "pdf_url":  _url(raw.get("pdf_path")),
        "docx_url": _url(raw.get("docx_path")),
        "pptx_url": _url(raw.get("pptx_path")),
        "xlsx_url": _url(raw.get("xlsx_path")),
    }


# ── UPLOAD ────────────────────────────────────────────────────────────────

@app.post("/api/upload")
async def api_upload(
    file: UploadFile = File(...),
    module: str = Form(default="upload"),
    db: Session = Depends(get_db),
    user: User | None = Depends(get_current_user_optional),
) -> dict:
    try:
        raw_bytes = await file.read()
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Could not read file: {exc}")
    filename     = file.filename or "uploaded_file"
    text_content = _extract_text(raw_bytes, filename)
    save_uploaded_material(filename=filename, module=module, extracted_text=text_content)
    summary_result = _upload.summarise(text_content, filename=filename)
    pdf_path  = export_pdf(summary_result["title"],  summary_result["sections"])
    docx_path = export_docx(summary_result["title"], summary_result["sections"])
    _save_to_db_and_search(db, user.id if user else None, f"[Upload:{filename}]",
                           "upload", summary_result["title"], summary_result["body"],
                           summary_result["sections"], {"pdf_path": pdf_path, "docx_path": docx_path})
    return {
        "success":    True,
        "filename":   filename,
        "module":     "upload",
        "module_name":"Upload AI",
        "title":      summary_result["title"],
        "body":       summary_result["body"],
        "sections":   summary_result["sections"],
        "keywords":   summary_result.get("keywords", []),
        "word_count": summary_result.get("word_count", 0),
        "pdf_url":    _url(pdf_path),
        "docx_url":   _url(docx_path),
    }


# ── WORKSPACE / DOCUMENTS ────────────────────────────────────────────────

@app.get("/workspace/history")
def workspace_history(
    db: Session = Depends(get_db),
    user: User | None = Depends(get_current_user_optional),
) -> list:
    if user:
        docs = list_documents(db, user_id=user.id)
        return [
            {
                "id":       d.id,
                "title":    d.title,
                "module":   d.module,
                "prompt":   d.prompt[:150],
                "preview":  d.body[:300],
                "pdf_url":  d.pdf_url,
                "docx_url": d.docx_url,
                "pptx_url": d.pptx_url,
                "xlsx_url": d.xlsx_url,
                "timestamp": d.created_at.isoformat() if d.created_at else None,
            }
            for d in docs
        ]
    return list_workspace_items()

@app.delete("/workspace/clear")
def workspace_clear(db: Session = Depends(get_db),
                    user: User | None = Depends(get_current_user_optional)) -> dict:
    if user:
        for d in db.query(Document).filter(Document.user_id == user.id).all():
            d.is_deleted = True
        db.commit()
    else:
        from workspace_store import clear_workspace
        try:
            clear_workspace()
        except Exception:
            pass
    return {"cleared": True}


# ── DOCUMENT SHARE ────────────────────────────────────────────────────────

@app.post("/api/document/{doc_id}/share")
def share_document(
    doc_id: str,
    mode: str = "view",
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> dict:
    doc = db.query(Document).filter(Document.id == doc_id, Document.user_id == user.id).first()
    if not doc:
        raise HTTPException(status_code=404, detail="Document not found.")
    if not doc.share_token:
        doc.share_token = secrets.token_urlsafe(32)
        doc.share_mode  = mode
        db.commit()
    return {"share_token": doc.share_token, "share_mode": doc.share_mode,
            "share_url": f"/shared/{doc.share_token}"}

@app.get("/shared/{token}")
def view_shared(token: str, db: Session = Depends(get_db)) -> dict:
    doc = get_document_by_share_token(db, token)
    if not doc:
        raise HTTPException(status_code=404, detail="Shared document not found.")
    secs = json.loads(doc.sections_json) if doc.sections_json else []
    return {"title": doc.title, "module": doc.module, "sections": secs,
            "body": doc.body, "share_mode": doc.share_mode}


# ── USER TEMPLATES ────────────────────────────────────────────────────────

class SaveTemplateRequest(BaseModel):
    module: str
    label: str
    prompt: str

@app.get("/api/templates")
def get_templates(
    module: str = "",
    db: Session = Depends(get_db),
    user: User | None = Depends(get_current_user_optional),
) -> list:
    system = [t for t in TEMPLATES if not module or t["module"] == module]
    if user:
        user_tpls = list_user_templates(db, user.id)
        custom = [
            {"id": t.id, "module": t.module, "label": t.label, "prompt": t.prompt, "custom": True}
            for t in user_tpls
            if not module or t.module == module
        ]
        return system + custom
    return system

@app.post("/api/templates")
def save_template(
    req: SaveTemplateRequest,
    db: Session = Depends(get_db),
    user: User = Depends(get_current_user),
) -> dict:
    t = save_user_template(db, user.id, req.module, req.label, req.prompt)
    return {"id": t.id, "module": t.module, "label": t.label, "prompt": t.prompt, "custom": True}


# ── MODULES ───────────────────────────────────────────────────────────────

@app.get("/api/modules")
def list_modules() -> list:
    return [{"id": k, "name": v, "icon": _MODULE_ICON.get(k, "🤖")} for k, v in _MODULE_NAMES.items()]


# ── PUBLIC DEVELOPER API (v1) ─────────────────────────────────────────────

async def _get_api_key_user(
    x_api_key: str | None = None,
    db: Session = Depends(get_db),
) -> User:
    from fastapi import Header
    if not x_api_key:
        raise HTTPException(status_code=401, detail="API key required. Pass X-API-Key header.")
    user = get_api_key_owner(db, x_api_key)
    if not user:
        raise HTTPException(status_code=401, detail="Invalid or inactive API key.")
    return user

from fastapi import Header

@app.post("/api/v1/generate")
def v1_generate(
    req: BuildRequest,
    x_api_key: str | None = Header(default=None),
    db: Session = Depends(get_db),
) -> dict:
    """Public API endpoint for developers. Requires X-API-Key header."""
    if not x_api_key:
        raise HTTPException(status_code=401, detail="API key required. Pass X-API-Key header.")
    user = get_api_key_owner(db, x_api_key)
    if not user:
        raise HTTPException(status_code=401, detail="Invalid API key.")
    allowed, _ = check_rate_limit(db, user)
    if not allowed:
        raise HTTPException(status_code=429, detail="Rate limit exceeded.")
    try:
        raw = _brain.run(req.prompt.strip(), module_hint=req.module)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Generation failed: {e}")
    increment_gen_count(db, user)
    return {
        "success": True, "module": raw.get("module"), "title": raw["title"],
        "body": raw["body"], "sections": raw["sections"],
        "pdf_url": _url(raw.get("pdf_path")), "docx_url": _url(raw.get("docx_path")),
        "pptx_url": _url(raw.get("pptx_path")), "xlsx_url": _url(raw.get("xlsx_path")),
    }

@app.get("/api/v1/user/documents")
def v1_list_documents(
    x_api_key: str | None = Header(default=None),
    db: Session = Depends(get_db),
) -> list:
    if not x_api_key:
        raise HTTPException(status_code=401, detail="API key required.")
    user = get_api_key_owner(db, x_api_key)
    if not user:
        raise HTTPException(status_code=401, detail="Invalid API key.")
    docs = list_documents(db, user_id=user.id)
    return [{"id": d.id, "title": d.title, "module": d.module,
             "pdf_url": d.pdf_url, "created_at": d.created_at.isoformat() if d.created_at else None}
            for d in docs]


# ── API KEY MANAGEMENT ────────────────────────────────────────────────────

class CreateApiKeyRequest(BaseModel):
    name: str = "My API Key"

@app.get("/api/keys")
def list_api_keys(user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> list:
    keys = db.query(ApiKey).filter(ApiKey.user_id == user.id, ApiKey.is_active == True).all()
    return [{"id": k.id, "name": k.name, "key": k.key[:12] + "…",  # masked
             "request_count": k.request_count, "last_used": k.last_used.isoformat() if k.last_used else None,
             "created_at": k.created_at.isoformat() if k.created_at else None} for k in keys]

@app.post("/api/keys")
def create_key(req: CreateApiKeyRequest, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> dict:
    k = create_api_key(db, user.id, name=req.name)
    return {"id": k.id, "name": k.name, "key": k.key,  # full key shown once
            "message": "Save this key — it will not be shown again."}

@app.delete("/api/keys/{key_id}")
def delete_api_key(key_id: str, user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> dict:
    k = db.query(ApiKey).filter(ApiKey.id == key_id, ApiKey.user_id == user.id).first()
    if not k:
        raise HTTPException(status_code=404, detail="API key not found.")
    k.is_active = False
    db.commit()
    return {"deleted": True}


# ── EDUCATION ENDPOINTS ───────────────────────────────────────────────────

class EducationGenerateRequest(BaseModel):
    prompt: str
    module: str | None = None

class EducationChatRequest(BaseModel):
    message: str
    module: str | None = None

@app.get("/education/modules")
def list_education_modules() -> list:
    return MODULE_INFO

@app.post("/education/generate")
def education_generate(
    req: EducationGenerateRequest,
    db: Session = Depends(get_db),
    user: User | None = Depends(get_current_user_optional),
) -> dict:
    try:
        result = _edu.generate(req.prompt, module=req.module)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Education generation failed: {exc}")
    pdf_path  = export_pdf(result["title"],  result["sections"])
    docx_path = export_docx(result["title"], result["sections"])
    _save_to_db_and_search(db, user.id if user else None, req.prompt, result["module"],
                           result["title"], result["body"], result["sections"],
                           {"pdf_path": pdf_path, "docx_path": docx_path})
    return {
        "module":   result["module"],
        "title":    result["title"],
        "body":     result["body"],
        "sections": result["sections"],
        "pdf_url":  _url(pdf_path),
        "docx_url": _url(docx_path),
    }

@app.post("/education/chat")
def education_chat(req: EducationChatRequest) -> dict:
    module = req.module or "student"
    add_chat_message(module, "user", req.message)
    try:
        response = _edu.chat(req.message, module=module)
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Chat failed: {exc}")
    reply = response.get("reply") or str(response)
    add_chat_message(module, "assistant", reply)
    return {"module": module, "reply": reply, "history": get_chat_history(module)}

@app.post("/education/upload-material")
async def education_upload_material(
    file: UploadFile = File(...),
    module: str = "student",
    db: Session = Depends(get_db),
    user: User | None = Depends(get_current_user_optional),
) -> dict:
    raw = await file.read()
    text_content = _extract_text(raw, file.filename or "file")
    save_uploaded_material(filename=file.filename or "unknown", module=module, extracted_text=text_content)
    combined = f"Summarise and create a study guide for:\n\n{text_content[:4000]}"
    try:
        result = _edu.generate(combined, module=module or "student")
    except Exception as exc:
        raise HTTPException(status_code=500, detail=f"Upload processing failed: {exc}")
    pdf_path  = export_pdf(result["title"],  result["sections"])
    docx_path = export_docx(result["title"], result["sections"])
    return {
        "filename": file.filename, "module": result["module"],
        "title": result["title"], "body": result["body"], "sections": result["sections"],
        "pdf_url": _url(pdf_path), "docx_url": _url(docx_path),
    }

@app.get("/education/materials")
def list_materials() -> list:
    return list_uploaded_materials()


# ── ROLE-BASED DASHBOARD ─────────────────────────────────────────────────

_ROLE_MODULES = {
    "student":   ["student", "exam", "teacher", "course", "research", "document"],
    "teacher":   ["teacher", "exam", "student", "admin", "course", "simulation"],
    "professor": ["professor", "research", "exam", "integrity", "document", "analytics"],
    "lawyer":    ["court", "document", "simulation", "integrity", "research", "admin"],
    "admin":     ["admin", "analytics", "document", "business", "content", "research"],
}

@app.get("/api/dashboard")
def role_dashboard(user: User = Depends(get_current_user), db: Session = Depends(get_db)) -> dict:
    """Returns personalised dashboard data for the user's role."""
    modules = _ROLE_MODULES.get(user.role, list(_MODULE_NAMES.keys())[:6])
    recent_docs = list_documents(db, user_id=user.id, limit=5)
    return {
        "role":           user.role,
        "recommended_modules": [
            {"id": m, "name": _MODULE_NAMES.get(m, m), "icon": _MODULE_ICON.get(m, "🤖")}
            for m in modules
        ],
        "recent_documents": [
            {"id": d.id, "title": d.title, "module": d.module,
             "created_at": d.created_at.isoformat() if d.created_at else None}
            for d in recent_docs
        ],
        "stats": {
            "total_generated": user.gen_count,
            "plan": user.plan,
        }
    }


# ── SPA CATCH-ALL ─────────────────────────────────────────────────────────

@app.get("/{full_path:path}", response_class=HTMLResponse, response_model=None, include_in_schema=False)
def serve_spa(full_path: str):
    if _DIST_DIR.is_dir():
        candidate = _DIST_DIR / full_path
        if candidate.is_file():
            return FileResponse(str(candidate))
        index = _DIST_DIR / "index.html"
        if index.is_file():
            return FileResponse(str(index))
    return HTMLResponse("<h1>Frontend not built.</h1>", status_code=503)
